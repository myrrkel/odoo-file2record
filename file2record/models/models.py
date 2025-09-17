# Copyright (C) 2024 - Michel Perrocheau (https://github.com/myrrkel).
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl.html).

from odoo import models, _, api
from odoo.tools import plaintext2html, html2plaintext, html_sanitize
from odoo.exceptions import UserError
from odoo.tools.pdf import OdooPdfFileReader
import openpyxl
import csv
import fitz
import json
import mammoth
import logging
import base64
import io
import re
import lxml.html.clean as html_clean
_logger = logging.getLogger(__name__)


def get_pdf_text(content, drop_last_page=False):
    buffer = io.BytesIO(content)
    pdf_reader = OdooPdfFileReader(buffer, strict=False)
    if pdf_reader.pages:
        if len(pdf_reader.pages) > 1 and drop_last_page:
            pages = pdf_reader.pages[:-1]
        else:
            pages = pdf_reader.pages
        text_content = '\n'.join([page.extract_text() for page in pages])
        pdf_reader.stream.close()
        return text_content


def extract_xlsx_content(raw_content, output_format='markdown'):
    workbook = openpyxl.load_workbook(io.BytesIO(raw_content), data_only=True)
    content = []

    for i, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        sheet_content = [f"# Sheet {i + 1} : {sheet_name}\n"]
        if output_format == 'markdown':
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) if cell else '' for cell in row)
                sheet_content.append("| %s |" % row_text)
        elif output_format == 'csv':
            output = io.StringIO()
            csv_writer = csv.writer(output)
            csv_writer.writerows(sheet.iter_rows(values_only=True))
            sheet_content.append(output.getvalue())

        content.append("\n".join(sheet_content))

    return "\n".join(content)

EXCLUDED_REQUIRED_FIELDS = {
    'product.template': ['product_variant_ids'],
}


# noinspection PyInconsistentReturns
class BaseModel(models.AbstractModel):
    _inherit = 'base'

    @api.model
    def is_file_to_record_button_visible(self):
        if self.env['file2record.config'].sudo().search([('model', '=', self._name),
                                                  ('show_upload_file_button', '=', True)]):
            return True

    def _is_attachment_txt(self, attachment_id):
        return attachment_id.mimetype in ['text/plain', 'text/html']

    def _is_attachment_xml(self, attachment_id, content):
        return content.startswith('<?xml') or attachment_id.name.endswith('.xml')

    def _is_attachment_html(self, attachment_id):
        extension = attachment_id.name.lower().split('.')[-1]
        return extension in ['html', 'xhtml', '.htm']

    def _is_attachment_document(self, attachment_id):
        extension = attachment_id.name.lower().split('.')[-1]
        return 'word' in attachment_id.mimetype or extension in ['docx', 'odt']

    def _is_attachment_xls(self, attachment_id):
        extension = attachment_id.name.lower().split('.')[-1]
        return 'spreadsheet' in attachment_id.mimetype or extension in ['xlsx', 'xls']

    def _create_record_from_attachment(self, res_id, default_values=None):
        attachment_id = self.env['ir.attachment'].browse(res_id)
        values = self._get_values_from_attachment_id(res_id)
        if values:
            if default_values:
                values.update(default_values)
            if attachment_id.res_id:
                record_id = self.env[attachment_id.res_model].browse(attachment_id.res_id)
                record_id.write(values)
                return record_id
            else:
                record_id = self._create_record_from_dict(values)
                if record_id:
                    attachment_id.res_id = record_id.id
                    attachment_id.res_model = self._name
                    attachment_id.register_as_main_attachment()
                    return record_id

    def _get_values_from_attachment_id(self, attachment_id):
        attachment_id = self.env['ir.attachment'].browse(attachment_id)
        content = base64.b64decode(attachment_id.with_context(bin_size=False).datas)
        return self._get_values_from_attachment(attachment_id, content)

    def _get_values_from_attachment(self, attachment_id, content):
        if 'pdf' in attachment_id.mimetype and isinstance(content, bytes):
            values = self._get_record_values(attachment_id.name, 'pdf', content)
        elif self._is_attachment_document(attachment_id):
            values = self._get_record_values(attachment_id.name, 'doc', content)
        elif self._is_attachment_xls(attachment_id):
            values = self._get_record_values(attachment_id.name, 'xls', content)
        elif self._is_attachment_txt(attachment_id):
            if isinstance(content, bytes):
                content = content.decode()
            if self._is_attachment_xml(attachment_id, content):
                values = self._get_record_values(attachment_id.name, 'xml', content)
            elif self._is_attachment_html(attachment_id):
                values = self._get_record_values(attachment_id.name, 'html', content)
            else:
                values = self._get_record_values(attachment_id.name, 'text', content)
        else:
            values = {}
        return values

    def _get_record_values(self, name, content_type, content):
        if not content:
            return {}
        config_id = self.get_file2record_config(content_type)
        if config_id and config_id.record_creation_method == 'code':
            values = config_id.eval_record_creation_code(content)
        else:
            values = self._get_record_values_from_raw_content(name, content_type, content)
        if config_id and config_id.post_process == 'code':
            values = config_id.eval_post_process_code(values)
        if config_id and config_id.post_process == 'method':
            post_process_function = getattr(self.env[config_id.model], config_id.model_post_process_method)
            values = post_process_function(values)

        if not values:
            values = {'name': name.split('.')[0]}
        if not values.get('name'):
            values['name'] = name.split('.')[0]

        return values

    def _get_record_values_from_raw_content(self, name, content_type, raw_content):
        content = ''
        if isinstance(raw_content, bytes):
            if content_type == 'pdf':
                content = self._get_html_from_pdf(raw_content)
            if content_type == 'doc':
                res = mammoth.convert_to_html(io.BytesIO(raw_content))
                content = self._clean_html(res.value)
            if content_type == 'xls':
                content = extract_xlsx_content(raw_content, output_format='csv')
        elif content_type == 'html':
            content = self._clean_html(raw_content)
        else:
            content = raw_content

        return self._get_record_values_from_content(name, content_type, content)

    def _get_record_values_from_content(self, name, content_type, content):
        values = {}
        if content_type == 'pdf':
            if isinstance(content, bytes):
                values = self._get_record_values_from_pdf(name, content)
            elif isinstance(content, str):
                values = self._get_record_values_from_text(name, content)
        if content_type == 'doc':
            values = self._get_record_values_from_doc(name, content)
        elif content_type == 'xml':
            values = self._get_record_values_from_xml(name, content)
        elif content_type == 'html':
            values = self._get_record_values_from_html(name, content)
        elif content_type == 'text':
            values = self._get_record_values_from_text(name, content)
        return values

    def _get_record_values_from_text(self, name, content):
        return {}

    def _get_record_values_from_xml(self, name, content):
        return self._get_record_values_from_text(name, content)

    def _get_record_values_from_doc(self, name, content):
        return self._get_record_values_from_html(name, content)

    def _get_record_values_from_html(self, name, content):
        res = self._get_record_values_from_text(name, content)
        return res

    def _get_record_values_from_pdf(self, name, content):
        return self._get_record_values_from_text(name, content)

    def _clean_html(self, html):
        p = re.compile(r'<img[^>]*>')
        html = p.sub('', html)
        try:
            html = html_clean.clean_html(html)
            res = html_sanitize(html, strip_style=True, strip_classes=True, sanitize_attributes=True, sanitize_style=True)
            return str(res)
        except Exception as err:
            _logger.error(err, exc_info=True)
            return html

    def _get_html_from_pdf(self, content, drop_last_page=False):
        doc = fitz.open("pdf", content)
        html_content = ''
        for i, page in enumerate(doc):
            if drop_last_page and i == len(doc) - 1:
                continue
            page.read_contents()
            img_list = page.get_images()
            for img in img_list:
                try:
                    page.delete_image(img[0])
                except Exception as err:
                    _logger.error(err, exc_info=True)
                    pass
            html_content += page.get_text('html', clip=None)
        return self._clean_html(html_content)

    def _exclude_fields_from_cleanup(self):
        return []

    def check_required_fields(self, values):
        def field_is_required(field):
            if self._name in EXCLUDED_REQUIRED_FIELDS:
                if field.name in EXCLUDED_REQUIRED_FIELDS[self._name]:
                    return False
            exclude_args = ['default', 'compute', 'company_dependent']
            if field.required and field.args and not any(arg in exclude_args for arg in field.args.keys()):
                return True

        required_fields = [f for f in self._fields if field_is_required(self._fields[f])]
        for required_field in required_fields:
            if not values.get(required_field):
                raise UserError("The field %s (%s) is required for %s" % (self._fields[required_field].string,
                                                                          required_field, self._name))
        return values

    def cleanup_record_values(self, values):
        if 'id' in values:
            values.pop('id')
        res = {key: values[key] for key in values.keys() if key in self._fields}
        for key in res.keys():
            value = res[key]
            if self._fields[key].type == 'integer':
                if isinstance(value, str):
                    value = re.sub("[^0-9]", "", value)
                    res[key] = int(value) if value else 0
            elif self._fields[key].type == 'float':
                if isinstance(value, str):
                    value = re.sub("[^0-9^,^.]", "", value).replace(',', '.')
                    res[key] = float(value) if value else 0
            elif self._fields[key].type == 'boolean':
                if isinstance(value, str):
                    val = value.lower()
                    if val in ('y', 'yes', 't', 'true', 'on', '1'):
                        value = True
                    elif val in ('n', 'no', 'f', 'false', 'off', '0'):
                        value = False
                    else:
                        value = None
                    res[key] = value
        return res

    def _get_default_record_creation_prompt(self, content, additional_instructions=''):
        json_dict = self._get_json_model_fields_description()
        instructions = '''You are an API. Return the {model} json dictionary provided
with values found in the document.
Don't do assertions.
If you don't find a value in the document, set the value at false.
If there is no relevant information in the document return an empty dictionary.'''.format(model=self._description)
        prompt_list = [('INSTRUCTIONS', instructions),
                       ('JSON DICTIONARY', json_dict),
                       ('FIELDS DESCRIPTION', self._get_fields_description(json.loads(json_dict))),
                       ('DOCUMENT', content)]
        if additional_instructions:
            prompt_list.insert(1, ('ADDITIONAL INSTRUCTIONS', additional_instructions))
        return '\n\n'.join('# %s :\n\n%s' % (key, value) for key, value in prompt_list)

    def get_env_file2record_config(self):
        if self.env.context.get('file2record_config_id'):
            file2record_config_id = self.env.context.get('file2record_config_id')
            return self.env['file2record.config'].browse(file2record_config_id)

    def model_description_excluded_fields(self):
        res = ['id', 'access_token', 'password', 'create_date', 'write_date']
        config_id = self.get_env_file2record_config()
        if config_id:
            res.extend(config_id.sudo().excluded_fields.mapped('name'))
        return res

    def model_description_fields(self):
        res = []
        config_id = self.get_env_file2record_config()
        if config_id:
            res.extend(config_id.sudo().fields.mapped('name'))
        return res

    def _get_model_fields(self):
        field_types = ['html', 'text', 'char', 'boolean', 'integer', 'date', 'datetime',
                       'float', 'many2one', 'one2many', 'monetary']
        excluded_fields = self.model_description_excluded_fields()
        fields = self.model_description_fields()

        def is_valid_field(field):
            if fields and field.name not in fields:
                return False
            if field.name in excluded_fields:
                return False
            if not field.store:
                return False
            if field.type not in field_types:
                return False
            if field.type == 'many2one':
                if not self.env[field.comodel_name]._get_json_model_many2one_field_description():
                    return False
            if field.type == 'one2many' and not field.comodel_name:
                if not self.env[field.comodel_name]._get_json_model_one2many_field_description():
                    return False
            return True
        model_fields = [self._fields[key] for key in self._fields if is_valid_field(self._fields[key])]
        return model_fields

    def _get_json_model_fields_description(self):
        model_fields = self._get_model_fields()
        empty_dict = {field.name: '' for field in model_fields}
        for field in model_fields:
            if field.type == 'many2one':
                field_description = self.env[field.comodel_name]._get_json_model_many2one_field_description()
                if field_description:
                    empty_dict[field.name] = field_description
                else:
                    empty_dict.pop(field.name)
            elif field.type == 'one2many':
                field_description = self.env[field.comodel_name]._get_json_model_one2many_field_description()
                if field_description:
                    empty_dict[field.name] = [field_description]
                else:
                    empty_dict.pop(field.name)
            elif field.type in ['integer', 'float', 'monetary']:
                empty_dict[field.name] = 0
            elif field.type in ['boolean']:
                empty_dict[field.name] = False
            elif field.type in ['date', 'datetime']:
                empty_dict[field.name] = '2000-01-01'
            elif field.type in ['datetime']:
                empty_dict[field.name] = '2000-01-01 00:00:00'
            else:
                empty_dict[field.name] = ''
        description = json.dumps(empty_dict, indent=2)
        return description

    def _get_field_description(self, field):
        help_text = '%s,%s' % (field.string, field.help) if field.help else field.string
        return '%s (%s) : %s' % (field.name, field.type, help_text)

    def _get_fields_description(self, json_values):
        fields = self._get_model_fields()
        fields = [field for field in fields if field.name in json_values]

        def _field_description(field):
            field_description = self._get_field_description(field)
            if field.type == 'many2one':
                many2one_description = self.env[field.comodel_name]._get_fields_description(json_values[field.name])
                if many2one_description and many2one_description not in descriptions:
                    descriptions.append(many2one_description)
            if field.type == 'one2many':
                one2many_description = self.env[field.comodel_name]._get_fields_description(json_values[field.name][0])
                if one2many_description and one2many_description not in descriptions:
                    descriptions.append(one2many_description)
            return field_description

        descriptions = []
        for field in fields:
            description = _field_description(field)
            if description not in descriptions:
                descriptions.append(description)

        fields_description = '\n'.join(descriptions)
        return fields_description
    
    def _get_one2many_field_description(self):
        return {}
    
    def _get_json_model_many2one_field_description(self):
        return {}

    def _get_json_model_one2many_field_description(self):
        return {}

    def _find_or_create_many2one_domain(self, values):
        return []

    def _find_or_create_many2one_record(self, values):
        if isinstance(values, int):
            return values
        domain = self._find_or_create_many2one_domain(values)
        if domain:
            try:
                res = self.search(domain, limit=1)
                return res[0].id if res else False
            except Exception as e:
                _logger.info(e)
                return False

    def _create_one2many_record(self, values_list):
        res = []
        for values in values_list:
            values = self._find_or_create_related_records(values)
            # values = self.cleanup_record_values(values)
            res.append(values)
        return res

    def _find_or_create_related_records(self, values):
        keys = [str(key) for key in values.keys()]
        for key in keys:
            if key not in self._fields:
                continue
            if self._fields[key].type == 'many2one':
                rec_id = False
                if values[key]:
                    model = self._fields[key].comodel_name
                    rec_id = self.env[model].with_context(field_name=key)._find_or_create_many2one_record(values[key])
                if rec_id:
                    values[key] = rec_id
                else:
                    values.pop(key)
            if self._fields[key].type == 'one2many':
                if values[key]:
                    model = self._fields[key].comodel_name
                    one2many_values_list = self.env[model].with_context(field_name=key)._create_one2many_record(
                        values[key])
                    if one2many_values_list:
                        one2many_values_list = [self.env[model].cleanup_record_values(v) for v in one2many_values_list]
                    if one2many_values_list:
                        values[key] = [(0, 0, one2many_values) for one2many_values in one2many_values_list]
                    else:
                        values.pop(key)
                else:
                    values.pop(key)
        return values

    def _create_record_from_dict(self, values):
        try:
            _logger.info('Raw values: %s', values)
            values = self._find_or_create_related_records(values)
            _logger.info('Create values: %s', values)
            values = self.cleanup_record_values(values)
            _logger.info('Cleaned up values: %s', values)
            values = self.check_required_fields(values)
            _logger.info('Cleaned up values: %s', values)
            return self.env[self._name].create(values)
        except Exception as err:
            _logger.error(err, exc_info=True, stack_info=True)
            raise err

    @api.model
    def create_records_from_attachments(self, res_ids, default_values=None):
        res = []
        context = {'lang': self.env.user.lang}
        for res_id in res_ids:
            try:
                record_id = self.with_context(context)._create_record_from_attachment(res_id, default_values)
                if record_id:
                    res.append(record_id.id)
            except Exception as err:
                _logger.error(err, exc_info=True)
                return self.get_upload_error_action()
        if not res:
            return self.get_upload_error_action()
        return self.get_record_creation_result_action(res)

    def get_upload_error_action(self):
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'message': _("Could not extract data from file."),
                'sticky': False,
                'type': 'danger',
            }
        }
    def get_record_creation_result_action(self, res):
        action_vals = {
            'name': _('Generated Records'),
            'domain': [('id', 'in', res)],
            'res_model': self._name,
            'type': 'ir.actions.act_window',
            'context': self._context
        }
        if len(res) == 1:
            action_vals.update({
                'views': [[False, "form"]],
                'view_mode': 'form',
                'res_id': res[0],
            })
        else:
            action_vals.update({
                'views': [[False, "list"], [False, "kanban"], [False, "form"]],
                'view_mode': 'list, kanban, form',
            })
        return action_vals

    def get_file2record_config(self, data_type):
        domain = [('model', '=', self._name),
                  ('data_type', '=', data_type)]
        config_id = self.env['file2record.config'].search(domain, limit=1)
        if not config_id:
            domain = [('model', '=', self._name),
                      ('data_type', '=', False)]
            config_id = self.env['file2record.config'].search(domain, limit=1)
        return config_id

    def get_attachments_without_record(self):
        domain = [('res_model', '=', self._name),
                  ('res_id', '=', False)]
        res = self.env['ir.attachment'].sudo().search(domain)
        return res.ids if res else []
